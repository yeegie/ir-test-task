import os

import pandas as pd
from openpyxl import Workbook


def export_result_table(
    code: str,
    barcode: str,
    table_data: pd.DataFrame,
    export_dir: str
) -> str:
    """
    Process and export data, return path to exported file.
    """
    filename = f"code_{code}.xlsx"
    output_path = os.path.join(export_dir, filename)

    # Унифицируем все коды в строки и убираем NaN
    codes = table_data["Коды"].dropna().astype(str).tolist()

    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"

    ws.cell(row=1, column=1, value="коды")
    ws.cell(row=2, column=1, value=str(barcode))

    for i, code_item in enumerate(codes, start=3):
        ws.cell(row=i, column=1, value=code_item)

    wb.save(output_path)
    return output_path